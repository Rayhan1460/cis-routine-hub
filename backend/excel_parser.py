import openpyxl
import re

def parse_excel(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    # 1. Read all data into a grid, resolving merged cells
    max_row = ws.max_row
    max_col = ws.max_column
    
    grid = [[None for _ in range(max_col)] for _ in range(max_row)]
    
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            grid[r-1][c-1] = ws.cell(row=r, column=c).value
            
    for merged_cell_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row_merge = merged_cell_range.bounds
        top_left_value = grid[min_row-1][min_col-1]
        for r in range(min_row, max_row_merge + 1):
            for c in range(min_col, max_col + 1):
                grid[r-1][c-1] = top_left_value

    days_of_week = ["Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    
    current_day = None
    time_slots = []
    
    classes = []
    warnings = []
    
    batches_set = set()
    sections_set = set()
    
    for row_idx, row in enumerate(grid):
        if not any(row):
            continue
            
        first_cell = str(row[1]).strip() if row[1] else ""
        
        # 2. DAY BLOCK DETECTION
        if first_cell in days_of_week:
            current_day = first_cell
            time_slots = []
            continue
            
        # 3. TIME SLOT DETECTION
        # Look for headers like Room, Course, Teacher to establish the time slots correctly
        # The row *above* this usually has the time ranges.
        if current_day and first_cell == 'Room':
            prev_row = grid[row_idx - 1]
            time_slots = []
            
            current_time = None
            for c_idx in range(2, len(row)):
                # Keep finding the time string from the row above
                if prev_row[c_idx] and isinstance(prev_row[c_idx], str) and '-' in prev_row[c_idx]:
                    current_time = str(prev_row[c_idx]).strip()
                elif prev_row[c_idx - 1] and isinstance(prev_row[c_idx - 1], str) and '-' in prev_row[c_idx - 1]:
                    current_time = str(prev_row[c_idx - 1]).strip()
                     
                if row[c_idx] == 'Course':
                    # Find matching teacher col for this course col
                    teacher_col = c_idx + 1 if (c_idx + 1 < len(row) and row[c_idx+1] == 'Teacher') else None
                    if current_time:
                        time_slots.append({
                            'time_str': current_time,
                            'course_col': c_idx,
                            'teacher_col': teacher_col
                        })
            continue
            
        # 4. DATA ROWS (Room detection)
        if current_day and time_slots:
            room = row[1]
            if not room or str(room).strip() == "":
                continue
                
            for ts in time_slots:
                course_val = row[ts['course_col']]
                teacher_val = row[ts['teacher_col']] if ts['teacher_col'] else None
                
                if course_val:
                    course_str = str(course_val).strip()
                    if not course_str:
                        continue
                        
                    # 5. BATCH AND SECTION EXTRACTION
                    batch = ""
                    section = ""
                    course_name = course_str
                    
                    match = re.match(r'^BATCH\s+(\d+)([A-Za-z0-9]+)\s+(.+)', course_str, re.IGNORECASE)
                    if match:
                        batch = match.group(1).upper()
                        section = match.group(2).upper()
                        course_name = match.group(3).strip()
                    else:
                        # For special cells like 'Makeup class', 'Project', preserve as is
                        batch = "N/A"
                        section = "N/A"
                        warnings.append(f"Non-standard course string preserved: {course_str}")
                        
                    if batch != "N/A":
                        batches_set.add(batch)
                        sections_set.add(section)
                        
                    # Time splitting
                    times = ts['time_str'].split('-')
                    start_time = times[0].strip() if len(times) > 0 else ""
                    end_time = times[1].strip() if len(times) > 1 else ""
                    
                    if start_time and "AM" not in start_time.upper() and "PM" not in start_time.upper():
                         if "AM" in end_time.upper(): start_time += " AM"
                         elif "PM" in end_time.upper():
                             if start_time.startswith("12") or start_time.startswith("01") or start_time.startswith("02") or start_time.startswith("03") or start_time.startswith("04") or start_time.startswith("05"):
                                 start_time += " PM"
                             else:
                                 start_time += " AM"
                                 
                    # 7. NORMALIZED ROUTINE OBJECT
                    classes.append({
                        'batch': batch,
                        'section': section,
                        'day': current_day,
                        'start_time': start_time,
                        'end_time': end_time,
                        'course_code': '', 
                        'course_name': course_name,
                        'teacher': str(teacher_val).strip() if teacher_val else "",
                        'room': str(room).strip()
                    })

    # Remove duplicates
    unique_classes = []
    seen = set()
    for c in classes:
        key = (c['batch'], c['section'], c['day'], c['start_time'], c['course_name'], c['room'])
        if key not in seen:
            seen.add(key)
            unique_classes.append(c)

    return {
        "batches": list(batches_set),
        "sections": list(sections_set),
        "classes": unique_classes,
        "warnings": warnings,
        "duplicates_removed": len(classes) - len(unique_classes)
    }

if __name__ == '__main__':
    res = parse_excel(r'C:\Users\hp\Downloads\ROUTINE_CIS_Fall 2026_VERSION 02 (Update). xlsx.xlsx')
    print("Batches:", len(res['batches']))
    print("Sections:", len(res['sections']))
    print("Classes:", len(res['classes']))
    print("Warnings:", len(res['warnings']))
    print("Duplicates:", res['duplicates_removed'])
